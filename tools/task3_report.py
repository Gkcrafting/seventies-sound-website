"""Add the Task 3 build and testing evidence to the shared report."""

import json
from pathlib import Path

from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TASK = ROOT / "website-testing"


def add_task3(doc):
    if any(
        paragraph.style.name == "Heading 1" and paragraph.text.startswith("Task 3:")
        for paragraph in doc.paragraphs
    ):
        raise ValueError("Task 3 is already in this document")

    def p(text):
        return doc.add_paragraph(text)

    def h(text):
        doc.add_heading(text, 2)

    def pic(path, caption, width):
        doc.add_picture(str(path), width=Inches(width))
        doc.inline_shapes[-1]._inline.docPr.set("descr", caption)
        doc.add_paragraph(caption, "Caption")

    doc.add_page_break()
    doc.add_heading("Task 3: Building and testing", 1)
    p(
        "I made the three pages in HTML, CSS and Javascript. I used one CSS file so the colours, fonts and links stay the same on each page. I kept the writing as normal webpage text, so it can be searched and read properly."
    )
    pic(
        ROOT / "website-design" / "website" / "screenshots" / "explore-desktop.png",
        "Explore page after building: search, genre choices, artist sections and listening links.",
        5.7,
    )
    p(
        "I kept all the artists on Explore because it is easier to look through them in one place. Search takes you there too. The genre buttons clear the search and show the selected genre. The request link checks the required boxes before it opens an email."
    )
    h("Changes made while building")
    p(
        "I made the Explore button open a list of the music styles. On a phone, the Menu button puts the page links into one list. I also added a Skip to content link and made the keyboard outline easier to see. The photo popup has a Close button, works with Escape and sends you back to the same button when it closes."
    )
    p(
        "The song videos only load after someone presses Play, so the page is not trying to load all of them straight away. I had a problem with YouTube in my first preview, so I used the localhost address in Live Server and it worked. Sources and credits open in a new tab, which means people do not lose their place on the site."
    )
    h("Functionality checks")
    p(
        "I tried the pages on Edge on my laptop and at phone size. I also used the keyboard for the menu, Skip link and bigger-photo button. I checked the YouTube videos in my localhost preview as well. My results are in website-testing/test-results/functional-results.json.")
    rows = json.loads((TASK / "test-results" / "functional-results.json").read_text())
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Check", "Result", "What happened")):
        cell.text = text
    for row in rows:
        line = table.add_row().cells
        line[0].text = row["test"]
        line[1].text = row["result"]
        line[2].text = row["note"]
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    p(
        "Everything above worked. I also played each of the seven videos in localhost. They all started and I did not get any script errors. That check is saved in website-design/draft-review/embed-checks.json."
    )
    h("Usability testing")
    p(
        "I made a short online feedback form so I can get opinions from classmates, friends, family and my teacher. I want people with different levels of interest in 1970s music to try the website, so I can see if it works for beginners as well as people who already know the songs."
    )
    p(
        "The form asks people to find music, use the search and genre buttons, open extra information, view a larger photo, try a video and use the Request page. It asks what was easy, what was confusing, how the site worked on their device and what should be improved. Name and age are optional, and the form does not ask for email addresses or other private information."
    )
    p(
        "I will look for repeated comments and ratings instead of changing the site because of one answer. The questions and testing plan are saved in website-testing/usability-testing.md. The feedback form is available at https://6nwnojh1fz.zite.so."
    )
    h("Review")
    p(
        "The main thing I changed was putting the artist information and song together on Explore. Before that, it felt like there were two separate lists. I also checked that the menu and photo popup worked with a keyboard, not just with a mouse. The three pages now have the same layout and navigation, and the parts I tested worked."
    )

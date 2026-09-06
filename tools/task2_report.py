from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont
import csv, json, hashlib

ROOT = Path(__file__).resolve().parents[1]
TASK = ROOT / "website-design"
FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"


def make_sheets():
    for name in ("home", "explore", "request"):
        im = Image.new("RGB", (2000, 1710), "white")
        d = ImageDraw.Draw(im)
        for label, pos in [("Wireframe", (10, 12)), ("Page design", (1010, 12))]:
            d.text(pos, label, font=ImageFont.truetype(FONT, 30), fill="black")
        wire = Image.open(TASK / "page-plans" / f"{name}.png")
        wire.thumbnail((980, 900))
        im.paste(wire, (10, 60))
        desktop = Image.open(TASK / "website/screenshots" / f"{name}-desktop.png")
        desktop.thumbnail((980, 1500))
        im.paste(desktop, (1010, 60))
        phone = Image.open(TASK / "website/screenshots" / f"{name}-mobile.png")
        phone = phone.crop((0, 0, 390, min(1000, phone.height)))
        phone.thumbnail((240, 680))
        im.paste(phone, (25, 1000))
        d.multiline_text(
            (295, 1020),
            "Phone layout (top)\n\nFull phone design and wireframe\nare saved with the design files.\n\nOne column for the content.\nSearch stays visible.\nThe same three page links\nare inside the Menu button.",
            font=ImageFont.truetype(FONT, 27),
            fill="black",
            spacing=12,
        )
        im.save(TASK / "website/screenshots" / f"{name}-sheet.png")
    im = Image.new("RGB", (1050, 1250), "white")
    d = ImageDraw.Draw(im)
    for name, x, label in [
        ("home-draft", 30, "First layout"),
        ("home", 550, "Revised layout"),
    ]:
        p = (
            TASK
            / ("draft-review" if name == "home-draft" else "website")
            / ("screenshots" if name == "home" else "")
            / f"{name}-mobile.png"
        )
        shot = Image.open(p).crop((0, 0, 390, 1100))
        d.text((x, 15), label, font=ImageFont.truetype(FONT, 30), fill="black")
        im.paste(shot, (x, 70))
    im.save(TASK / "draft-review/home-comparison.png")


def asset_rows():
    rows = []
    photos = {
        v["file"]: v
        for v in json.loads((TASK / "assets/licences/photos.json").read_text()).values()
    }
    for p in sorted((TASK / "assets").rglob("*")):
        if not p.is_file() or p.name in ("production.json", "asset-list.csv"):
            continue
        rel = p.relative_to(TASK / "assets").as_posix()
        if p.name in photos:
            item = photos[p.name]
            source = item["source"]
            rights = f"Photo: {item['author']}; {item['licence']}; {item['url']}"
            use = "Artist photo on Home / Explore"
            prep = "Downloaded unchanged; displayed at smaller sizes"
        elif p.name == "classic-tracks.json":
            source = "Official artist YouTube uploads and song sources in assets/written-content/classic-tracks.json"
            rights = "Copyright retained by the music/video rights holders. External players and links only."
            use = "70s listening choices and songwriter credits"
            prep = "Checked artist, year and official upload; kept recordings external"
        elif p.name in ("photos.txt", "photos.json"):
            source = "Wikimedia Commons file pages listed in this file"
            rights = "Photo credits and individual Creative Commons licences"
            use = "Image source and copyright records"
            prep = "Recorded author, source, licence and changes"
        elif p.parent.name == "icons":
            orig = "x" if p.stem == "close" else p.stem
            source = f"https://raw.githubusercontent.com/lucide-icons/lucide/main/icons/{orig}.svg"
            rights = "Lucide ISC; Feather-derived icon MIT. Full notices in licences/lucide.txt."
            use = "Navigation/search/dialog control; paired with a text label."
            prep = "Downloaded SVG; kept original paths. close.svg renamed from x.svg."
        elif p.parent.name == "licences":
            source = "https://github.com/lucide-icons/lucide/blob/main/LICENSE"
            rights = "Retain full notice"
            use = "Attribution"
            prep = "Unchanged"
        elif p.suffix == ".mp4":
            source = "Original project script and title cards; ElevenLabs Alice narration"
            rights = (
                "Original visuals/text; AI-generated narration, not a student recording"
            )
            use = "Home introduction video"
            prep = "H.264 720p, AAC audio and British narration; prepared for streaming"
        elif p.suffix == ".vtt":
            source = "Original video script"
            rights = "Original project text"
            use = "English captions"
            prep = "Timed caption file"
        elif p.name == "video-text.txt":
            source = "Original video script"
            rights = "Original project text"
            use = "Written video text beside player"
            prep = "Plain text"
        elif p.name == "page-copy.json":
            source = "Original wording based on the linked artist sources in this file"
            rights = "Original project text; facts attributed to sources"
            use = "Artist introductions and listening notes"
            prep = "Stored with source links"
        elif p.name == "poster.jpg":
            source = "Original first video title card"
            rights = "Original project artwork"
            use = "Video poster"
            prep = "1280 x 720 JPEG, quality 88"
        else:
            source = "Original project artwork"
            rights = "Original project artwork"
            use = (
                "Home genre cards / Explore image enlargement"
                if "record" in p.name
                else "Shared site identity"
            )
            prep = (
                "Editable SVG"
                if p.suffix == ".svg"
                else "800 x 480 WebP, quality 85; optional raster fallback"
            )
        rows.append(
            [
                rel,
                source,
                rights,
                use,
                prep,
                p.stat().st_size,
                hashlib.sha256(p.read_bytes()).hexdigest(),
            ]
        )
    with (TASK / "assets/asset-list.csv").open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(
            ["file", "source", "rights", "use", "preparation", "bytes", "sha256"]
        )
        w.writerows(rows)
    return rows


def add_task2(doc):
    if any(
        p.style.name == "Heading 1" and p.text.startswith("Task 2:")
        for p in doc.paragraphs
    ):
        raise ValueError("Task 2 is already in this document")

    def p(text):
        return doc.add_paragraph(text)

    def pic(path, caption, width=6.6):
        doc.add_picture(str(path), width=Inches(width))
        doc.inline_shapes[-1]._inline.docPr.set("descr", caption)
        doc.add_paragraph(caption, "Caption")

    def h(text):
        doc.add_heading(text, 2)

    doc.add_page_break()
    doc.add_heading("Task 2: Designs and assets", 1)
    p(
        "I kept the three pages from my plan. Georgia is used for headings and Arial for the main text, because both are easy to read. I used the same dark blue, gold and pale background on every page so it feels like one site."
    )
    pic(
        TASK / "website/screenshots/style-sheet.png",
        "Colours and text styles used across the website.", 4.8,
    )
    h("Final page designs")
    p("The three designs below show the wireframe, finished desktop page and phone layout. Full-size versions are saved in website-design/page-plans and website-design/website.")
    designs = doc.add_table(rows=1, cols=3)
    designs.style = "Table Grid"
    for cell, name, title in zip(designs.rows[0].cells, ("home", "explore", "request"), ("Home", "Explore", "Request")):
        cell.text = title
        run = cell.add_paragraph().add_run()
        run.add_picture(str(TASK / "website" / "screenshots" / f"{name}-sheet.png"), width=Inches(1.9))
    p("Home gives a quick way into the music. Explore keeps each artist, photo and song together, with extra facts hidden in accordions. Request uses a short labelled form. On a phone, the content becomes one column and the page links go inside Menu.")
    h("Important states")
    pic(
        TASK / "website/screenshots/states-desktop.png",
        "Drop-down, search, image popup and form messages.", 3.8,
    )
    p(
        "The Explore menu links to each genre. Search stays on Explore. A blank search gives help and no matches gives a way back to all artists. The image popup has a Close button and keyboard support. The request button checks the required boxes, then opens a pre-filled email for the visitor to send."
    )
    h("Review and assets")
    pic(
        TASK / "draft-review/home-comparison.png",
        "First Home phone layout compared with the revised layout.", 3.2,
    )
    p(
        "In my first Home design, the heading and record took up too much room, so the genre cards were too far down on a phone. I made the heading smaller and removed the fixed height. The cards now appear much sooner."
    )
    p(
        "The pages fitted both screen sizes and the links stayed in the same place. I also checked the popup with a keyboard. Explore has seven artists from different styles, and videos only load when someone presses Play."
    )
    p(
        "The logo and record illustrations were made in Canva, then saved as SVG files for the website. I used credited artist photos, Lucide icons and official YouTube embeds. The video has original title cards and script, an ElevenLabs voiceover, captions and a written transcript. Full source, licence, file-size and preparation details are in website-design/assets/asset-list.csv and licences/photos.txt. [13]"
    )
    p("[13] Lucide licence: https://lucide.dev/license. Accessed 5 September 2026.")

if __name__ == "__main__":
    make_sheets()
    asset_rows()
    path = ROOT / "report/Website development report.docx"
    doc = Document(path)
    before = [p.text for p in doc.paragraphs]
    add_task2(doc)
    assert [p.text for p in doc.paragraphs][: len(before)] == before
    doc.save(path)
    print("Task 2 added to the existing report.")
